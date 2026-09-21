import streamlit as st
import streamlit.components.v1 as components
import pyrebase
import firebase_admin
from firebase_admin import credentials, firestore, storage, auth as admin_auth
from docx import Document
from docx.shared import Inches
import io
import zipfile
import requests
import uuid
import re
from datetime import timedelta
import math
import random
import string
import time

# ==========================================
# 1. Firebase 初始化與設定 (採用快取避免重複連線)
# ==========================================
firebaseConfig = {
    "apiKey": "AIzaSyA-A3dn0hgJm9JR2NGvijRc3AD2x3sRgng",
    "authDomain": "amooo-b2c35.firebaseapp.com",
    "databaseURL": "", 
    "projectId": "amooo-b2c35",
    "storageBucket": "amooo-b2c35.firebasestorage.app",
    "messagingSenderId": "858227626420",
    "appId": "1:858227626420:web:d924c8a857f8c2542662e1",
    "measurementId": "G-ZJZ3ZZP3VM"
}

@st.cache_resource
def init_firebase_connection():
    try:
        firebase = pyrebase.initialize_app(firebaseConfig)
        pyrebase_auth = firebase.auth()
    except:
        pyrebase_auth = None

    if not firebase_admin._apps:
        try:
            if "firebase_key" in st.secrets:
                cert_dict = dict(st.secrets["firebase_key"])
                cert_dict["private_key"] = cert_dict["private_key"].replace("\\n", "\n")
                cred = credentials.Certificate(cert_dict)
            else:
                cred = credentials.Certificate('firebase-key.json')
        except Exception:
            cred = credentials.Certificate('firebase-key.json')

        firebase_admin.initialize_app(cred, {
            'storageBucket': firebaseConfig['storageBucket']
        })
    return pyrebase_auth, firestore.client()

auth, db = init_firebase_connection()

# 設定全局共享工作區 ID
SHARED_WORKSPACE = "global_shared_workspace"
ADMIN_EMAILS = [
    "jay03wn@amoo.com",
    "ru@amoo.com"
]

AVATAR_OPTIONS = {
    "黃色探險家": "https://api.dicebear.com/9.x/bottts/svg?seed=Felix",
    "藍色小怪": "https://api.dicebear.com/9.x/bottts/svg?seed=Aneka",
    "開心星人": "https://api.dicebear.com/9.x/bottts/svg?seed=Mimi",
    "酷炫機甲": "https://api.dicebear.com/9.x/bottts/svg?seed=Jack",
    "呆萌方塊": "https://api.dicebear.com/9.x/bottts/svg?seed=Jude",
    "粉紅助手": "https://api.dicebear.com/9.x/bottts/svg?seed=Sara"
}

# ==========================================
# 2. 網頁基本設定 & 狀態管理
# ==========================================
st.set_page_config(page_title="題庫分類系統", layout="wide")

components.html(
    """
    <script>
    if (!window.parent.document.getElementById('amo-keyboard')) {
        const script = window.parent.document.createElement('script');
        script.id = 'amo-keyboard';
        script.innerHTML = `
            document.addEventListener('keydown', function(e) {
                if (e.key === 'ArrowLeft') {
                    const btns = Array.from(document.querySelectorAll('button'));
                    const btn = btns.find(el => el.innerText.includes('上一題'));
                    if (btn) btn.click();
                }
                if (e.key === 'ArrowRight') {
                    const btns = Array.from(document.querySelectorAll('button'));
                    const btn = btns.find(el => el.innerText.includes('下一題'));
                    if (btn) btn.click();
                }
            });
        `;
        window.parent.document.head.appendChild(script);
    }
    </script>
    """, height=0, width=0
)

if "user" not in st.session_state:
    st.session_state.user = None
if "questions" not in st.session_state:
    st.session_state.questions = []
if "loaded_page" not in st.session_state:
    st.session_state.loaded_page = None  # Lazy Loading 頁面標記
if "categories" not in st.session_state:
    st.session_state.categories = ["生藥", "中藥", "法規", "實務"]
if "files_meta" not in st.session_state:
    st.session_state.files_meta = []
if "active_file_id" not in st.session_state:
    st.session_state.active_file_id = None
if "current_page" not in st.session_state:
    st.session_state.current_page = "檔案管理"
if "expanded_folders" not in st.session_state:
    st.session_state.expanded_folders = set()
if "last_used_folder" not in st.session_state:
    st.session_state.last_used_folder = None
if "recent_folders" not in st.session_state:
    st.session_state.recent_folders = []
if "all_users" not in st.session_state:
    st.session_state.all_users = {}

# ==========================================
# 3. 雲端同步與處理模組 (Lazy Loading 版)
# ==========================================
def upload_image_to_storage(img_bytes):
    bucket = storage.bucket()
    image_id = str(uuid.uuid4())
    blob = bucket.blob(f"workspaces/{SHARED_WORKSPACE}/images/{image_id}.png")
    blob.upload_from_string(img_bytes, content_type='image/png')
    url = blob.generate_signed_url(expiration=timedelta(days=3650))
    return url

def sync_user_meta():
    db.collection("workspaces").document(SHARED_WORKSPACE).set({
        "categories": st.session_state.categories,
        "files_meta": st.session_state.files_meta,
        "active_file_id": st.session_state.active_file_id
    }, merge=True)

def load_user_profile():
    if not st.session_state.user: return
    uid = st.session_state.user['localId']
    email = st.session_state.user.get('email', '')
    doc = db.collection("users").document(uid).get()
    
    if doc.exists:
        st.session_state.user_profile = doc.to_dict()
        updates = {}
        defaults = {"is_setup_complete": True, "points": 0, "categorized_count": 0, "inventory": [], "notifications": []}
        for k, v in defaults.items():
            if k not in st.session_state.user_profile:
                st.session_state.user_profile[k] = v
                updates[k] = v
        if updates:
            db.collection("users").document(uid).update(updates)
    else:
        default_profile = {
            "uid": uid, "email": email, "nickname": email.split('@')[0] if email else "新用戶",
            "avatar_url": "", "points": 0, "categorized_count": 0, "inventory": [], "notifications": [], "is_setup_complete": False
        }
        db.collection("users").document(uid).set(default_profile)
        st.session_state.user_profile = default_profile

@st.cache_data(ttl=300)
def fetch_all_users_cached():
    docs = db.collection("users").stream()
    valid_users = {}
    for doc in docs:
        data = doc.to_dict()
        if data and data.get("email"):
            valid_users[doc.id] = data
    return valid_users

def load_all_users():
    st.session_state.all_users = fetch_all_users_cached()

def load_from_cloud():
    load_user_profile()
    
    workspace_doc = db.collection("workspaces").document(SHARED_WORKSPACE).get()
    if workspace_doc.exists:
        data = workspace_doc.to_dict()
        st.session_state.categories = data.get("categories", ["生藥", "中藥", "法規", "實務"])
        st.session_state.files_meta = data.get("files_meta", [])
        st.session_state.active_file_id = data.get("active_file_id", None)
        st.session_state.category_counts = data.get("category_counts", {})
    
    st.session_state.questions = []
    st.session_state.loaded_page = None

def save_questions_to_firestore(new_questions):
    questions_ref = db.collection("workspaces").document(SHARED_WORKSPACE).collection("questions")
    for i in range(0, len(new_questions), 400):
        chunk = new_questions[i:i + 400]
        batch = db.batch()
        for q in chunk:
            batch.set(questions_ref.document(q["q_id"]), q)
        batch.commit()

def update_single_question_category(q_id, category, year_info, is_doubt=None):
    uid = st.session_state.user['localId']
    q = next((x for x in st.session_state.questions if x['q_id'] == q_id), None)
    if not q: return

    old_category = q.get("category", "未分類")
    was_unclassified = (old_category == "未分類")
    now_unclassified = (category == "未分類")
    file_id = q.get("file_id")

    # 1. 為了 Lazy Loading 更新 files_meta 的分類計數，並即時更新統計地圖 (大幅減少數據庫讀取)
    if old_category != category:
        db.collection("workspaces").document(SHARED_WORKSPACE).update({
            f"category_counts.{old_category}": firestore.Increment(-1),
            f"category_counts.{category}": firestore.Increment(1)
        })
        if "category_counts" in st.session_state:
            st.session_state.category_counts[old_category] = max(0, st.session_state.category_counts.get(old_category, 1) - 1)
            st.session_state.category_counts[category] = st.session_state.category_counts.get(category, 0) + 1

    if was_unclassified and not now_unclassified:
        for f in st.session_state.files_meta:
            if f["file_id"] == file_id:
                f["categorized_qs"] = f.get("categorized_qs", 0) + 1
                break
    elif not was_unclassified and now_unclassified:
        for f in st.session_state.files_meta:
            if f["file_id"] == file_id:
                f["categorized_qs"] = max(0, f.get("categorized_qs", 0) - 1)
                break

    # 2. 處理點數與使用者紀錄
    update_data = {"category": category, "year_info": year_info}
    if is_doubt is not None:
        update_data["is_doubt"] = is_doubt

    if category != "未分類":
        update_data["categorizer_uid"] = uid
        q["categorizer_uid"] = uid

        if was_unclassified:
            db.collection("users").document(uid).update({"categorized_count": firestore.Increment(1)})
            if "user_profile" in st.session_state:
                st.session_state.user_profile["categorized_count"] = st.session_state.user_profile.get("categorized_count", 0) + 1

        if not q.get("is_point_awarded"):
            try:
                q_num = int(year_info.split('-')[-1])
                if 41 <= q_num <= 80:
                    update_data["is_point_awarded"] = True
                    q["is_point_awarded"] = True
                    db.collection("users").document(uid).update({"points": firestore.Increment(1)})
                    if "user_profile" in st.session_state:
                        st.session_state.user_profile["points"] = st.session_state.user_profile.get("points", 0) + 1
            except:
                pass
    else:
        update_data["categorizer_uid"] = None
        q["categorizer_uid"] = None

    # 3. 寫入雲端與更新本地記憶體
    db.collection("workspaces").document(SHARED_WORKSPACE).collection("questions").document(q_id).update(update_data)
    q["category"] = category
    q["year_info"] = year_info
    if is_doubt is not None:
        q["is_doubt"] = is_doubt

def delete_file_data(file_id):
    questions_ref = db.collection("workspaces").document(SHARED_WORKSPACE).collection("questions")
    q_docs = questions_ref.where("file_id", "==", file_id).stream()
    batch = db.batch()
    count = 0
    for doc in q_docs:
        batch.delete(doc.reference)
        count += 1
        if count >= 400:
            batch.commit()
            batch = db.batch()
            count = 0
    if count > 0: batch.commit()

    st.session_state.questions = []
    st.session_state.loaded_page = None
    st.session_state.files_meta = [f for f in st.session_state.files_meta if f.get("file_id") != file_id]
    if st.session_state.active_file_id == file_id:
        st.session_state.active_file_id = None
    sync_user_meta()

def parse_docx_with_images(file, file_id):
    doc = Document(file)
    parsed_q = []
    q_count = 0 
    for table in doc.tables:
        for row in table.rows:
            q_blocks = []
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    para_text = ""
                    for run in paragraph.runs:
                        para_text += run.text
                        drawings = run._element.xpath('.//w:drawing')
                        for drawing in drawings:
                            for blip in drawing.xpath('.//a:blip'):
                                rId = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                                if rId:
                                    img_part = doc.part.related_parts[rId]
                                    img_bytes = img_part.blob
                                    if para_text.strip():
                                        q_blocks.append({"type": "text", "content": para_text.strip()})
                                        para_text = ""
                                    img_url = upload_image_to_storage(img_bytes)
                                    q_blocks.append({"type": "image", "content": img_url})
                    if para_text.strip():
                        q_blocks.append({"type": "text", "content": para_text.strip()})
            
            text_only = "".join([b["content"] for b in q_blocks if b["type"]=="text"])
            if q_blocks and text_only not in [q.get('_raw_text', '') for q in parsed_q]:
                parsed_q.append({
                    "q_id": str(uuid.uuid4()), 
                    "file_id": file_id, 
                    "order_index": q_count,  
                    "blocks": q_blocks, 
                    "category": "未分類", 
                    "year_info": "", 
                    "_raw_text": text_only,
                    "is_doubt": False,
                    "is_point_awarded": False
                })
                q_count += 1
    return parsed_q

def login_ui():
    st.title("系統登入")
    st.info("提示：目前系統設定為所有登入帳號皆共享同一份題庫與分類進度。", icon=":material/lightbulb:")
    email = st.text_input("帳號 (Email)")
    password = st.text_input("密碼", type="password")
    remember_me = st.checkbox("記住我的登入狀態", value=True)
    
    if st.button("安全登入", icon=":material/login:"):
        try:
            user = auth.sign_in_with_email_and_password(email, password)
            st.session_state.user = user
            if remember_me:
                st.query_params["remember_token"] = user['refreshToken']
            load_from_cloud()
            st.rerun()
        except Exception as e:
            st.error("登入失敗，請檢查帳號密碼。")

def setup_ui():
    st.title("歡迎！請完成初次帳號設定")
    st.write("這是您第一次登入，為了系統安全與運作，請重設密碼並指定專屬暱稱。")
    
    with st.container(border=True):
        new_nickname = st.text_input("設定您的暱稱", placeholder="輸入喜歡的名稱...")
        new_pwd = st.text_input("設定新密碼 (至少 6 碼)", type="password")
        confirm_pwd = st.text_input("再次確認新密碼", type="password")
        
        if st.button("完成設定並進入系統", type="primary", icon=":material/done_all:"):
            if not new_nickname.strip():
                st.error("暱稱不可為空，請輸入您的暱稱。", icon=":material/error:")
                return
            if len(new_pwd) < 6:
                st.error("密碼長度太短，請至少輸入 6 碼。", icon=":material/error:")
                return
            if new_pwd != confirm_pwd:
                st.error("兩次輸入的密碼不一致，請重新確認。", icon=":material/error:")
                return
            
            try:
                uid = st.session_state.user['localId']
                admin_auth.update_user(uid, password=new_pwd)
                db.collection("users").document(uid).update({
                    "nickname": new_nickname.strip(),
                    "is_setup_complete": True
                })
                st.session_state.user_profile["nickname"] = new_nickname.strip()
                st.session_state.user_profile["is_setup_complete"] = True
                
                st.success("設定完成！正在為您重新導向...", icon=":material/check_circle:")
                time.sleep(1.5)
                st.rerun()
            except Exception as e:
                st.error(f"設定失敗，發生錯誤：{e}", icon=":material/error:")


def get_year_from_info(year_info):
    if not year_info: return 0
    try: return int(str(year_info).split('-')[0])
    except: return 0

def hierarchical_select(label, categories, key_prefix, default_cat=None, include_unclassified=False, col_layout=None):
    main_folders = list(dict.fromkeys([c.split("/")[0] for c in categories]))
    if include_unclassified:
        main_folders.insert(0, "未分類")
        
    default_main_idx = 0
    if default_cat:
        if default_cat == "未分類" and include_unclassified:
            default_main_idx = 0
        else:
            default_main = default_cat.split("/")[0]
            if default_main in main_folders:
                default_main_idx = main_folders.index(default_main)
    
    if col_layout is None:
        c1, c2 = st.columns(2)
    else:
        c1, c2 = col_layout[0], col_layout[1]
        
    selected_main = c1.selectbox(f"{label} (母資料夾)", main_folders, index=default_main_idx, key=f"{key_prefix}_main")
    
    if selected_main == "未分類":
        return "未分類"
        
    sub_folders = [c for c in categories if c == selected_main or c.startswith(selected_main + "/")]
    default_sub_idx = 0
    if default_cat and default_cat in sub_folders:
        default_sub_idx = sub_folders.index(default_cat)
        
    selected_sub = c2.selectbox(f"{label} (子資料夾與路徑)", sub_folders, index=default_sub_idx, key=f"{key_prefix}_sub")
    return selected_sub

def render_categorizer_info(cat_uid):
    if cat_uid and 'all_users' in st.session_state and cat_uid in st.session_state.all_users:
        prof = st.session_state.all_users[cat_uid]
        avatar = prof.get("avatar_url")
        name = prof.get("nickname", "未知")
        if avatar:
            st.markdown(f'<img src="{avatar}" width="24" height="24" style="border-radius:50%; vertical-align:middle; object-fit: cover;"> <span style="vertical-align:middle; color:gray; font-size: 0.9em;">分類者: **{name}**</span>', unsafe_allow_html=True)
        else:
            st.markdown(f':material/account_circle: <span style="vertical-align:middle; color:gray; font-size: 0.9em;">分類者: **{name}**</span>', unsafe_allow_html=True)

# ==========================================
# 4. 主系統介面
# ==========================================
def main_app():
    is_admin = (st.session_state.user.get('email') == ADMIN_EMAIL)
    current_uid = st.session_state.user['localId']
    
    st.sidebar.title("導覽列")
    
    PAGES = {
        "檔案管理": ":material/folder_managed:",
        "資料夾管理": ":material/snippet_folder:",
        "題庫分類作業": ":material/rule_folder:",
        "全站題目搜索與瀏覽": ":material/search:",
        "題目匯出": ":material/download:",
        "個人檔案": ":material/person:"
    }
    
    if st.session_state.current_page not in PAGES:
        st.session_state.current_page = "檔案管理"
        
    for p_name, p_icon in PAGES.items():
        if st.sidebar.button(p_name, icon=p_icon, use_container_width=True, type="primary" if st.session_state.current_page == p_name else "secondary"):
            if st.session_state.current_page == "題庫分類作業" and p_name != "題庫分類作業":
                if st.session_state.active_file_id:
                    for f in st.session_state.files_meta:
                        if f.get("file_id") == st.session_state.active_file_id:
                            f["locked_by"] = None
                            f["locked_at"] = None
                    sync_user_meta()
            st.session_state.current_page = p_name
            st.rerun()
            
    st.sidebar.divider()
    if st.sidebar.button("登出系統", icon=":material/logout:"):
        if st.session_state.active_file_id:
            for f in st.session_state.files_meta:
                if f.get("file_id") == st.session_state.active_file_id:
                    f["locked_by"] = None
                    f["locked_at"] = None
            sync_user_meta()
            
        st.session_state.user = None
        st.query_params.clear()
        st.rerun()

    # ---------- 頁面 1：檔案管理 ----------
    if st.session_state.current_page == "檔案管理":
        c_head1, c_head2, c_head3 = st.columns([2, 1, 1], vertical_alignment="bottom")
        with c_head1:
            st.header("檔案管理")
        with c_head2:
            if st.button("立即刷新", icon=":material/refresh:", use_container_width=True):
                load_from_cloud()
                st.rerun()
        with c_head3:
            auto_refresh = st.toggle("自動更新狀態", value=False)
        
        if is_admin:
            with st.container(border=True):
                st.subheader("上傳新檔案", anchor=False)
                c1, c2 = st.columns(2)
                exam_year = c1.number_input("考試年份", 90, 150, 114)
                exam_session = c2.number_input("考試次數", 1, 10, 2)
                uploaded_file = st.file_uploader("選擇 Word 題庫檔", type=["docx"])
                
                if st.button("上傳並解析", icon=":material/cloud_upload:", type="primary"):
                    if uploaded_file:
                        with st.spinner("正在解析與上傳..."):
                            file_id = str(uuid.uuid4())
                            new_qs = parse_docx_with_images(uploaded_file, file_id)
                            save_questions_to_firestore(new_qs)
                            
                            st.session_state.files_meta.append({
                                "file_id": file_id, "name": uploaded_file.name,
                                "year": exam_year, "session": exam_session, "completed": False, "current_index": 0,
                                "locked_by": None, "locked_at": None,
                                "total_qs": len(new_qs),
                                "categorized_qs": 0
                            })
                            st.session_state.active_file_id = file_id
                            sync_user_meta()
                        st.toast("上傳完成")
                        st.rerun()

        st.markdown("### 已上傳檔案")
        if not st.session_state.files_meta:
            st.caption("目前系統尚未上傳任何檔案。")
            
        for f in st.session_state.files_meta:
            f_id = f["file_id"]
            
            # [優化] 若舊資料沒有題數緩存，自動補齊
            if "total_qs" not in f:
                q_docs = list(db.collection("workspaces").document(SHARED_WORKSPACE).collection("questions").where("file_id", "==", f_id).select(["category"]).stream())
                f["total_qs"] = len(q_docs)
                f["categorized_qs"] = len([doc for doc in q_docs if doc.to_dict().get("category", "未分類") != "未分類"])
                sync_user_meta()
                
            total_qs = f.get("total_qs", 0)
            categorized_qs = f.get("categorized_qs", 0)
            
            with st.container(border=True):
                col1, col2, col3, col4 = st.columns([1, 2.5, 1.2, 1], vertical_alignment="center")
                with col1:
                    is_completed = st.checkbox("標示完成", value=f.get("completed", False), key=f"chk_{f_id}")
                    if is_completed != f.get("completed", False):
                        f["completed"] = is_completed
                        sync_user_meta()
                with col2:
                    st.markdown(f"**{f['name']}**")
                    st.caption(f"考試資訊：{f.get('year', 114)} 年第 {f.get('session', 2)} 次")
                    if total_qs > 0:
                        st.progress(categorized_qs / total_qs, text=f"進度: {categorized_qs} / {total_qs} 題")
                with col3:
                    now = time.time()
                    locked_by = f.get("locked_by")
                    locked_at = f.get("locked_at", 0)
                    
                    if locked_by and (now - locked_at > 180):
                        locked_by = None
                        f["locked_by"] = None
                        f["locked_at"] = None
                        
                    if locked_by and locked_by != current_uid:
                        locker_prof = st.session_state.all_users.get(locked_by, {})
                        avatar = locker_prof.get("avatar_url")
                        name = locker_prof.get("nickname", "未知")
                        
                        st.markdown('<div style="text-align: center; margin-bottom: 5px;">', unsafe_allow_html=True)
                        if avatar:
                            st.markdown(f'<img src="{avatar}" width="20" height="20" style="border-radius:50%; vertical-align:middle; object-fit:cover;"> <span style="color:#d32f2f; font-size:0.85em;">**{name}** 編輯中</span>', unsafe_allow_html=True)
                        else:
                            st.markdown(f':material/account_circle: <span style="color:#d32f2f; font-size:0.85em;">**{name}** 編輯中</span>', unsafe_allow_html=True)
                        st.markdown('</div>', unsafe_allow_html=True)
                        
                        st.button("鎖定中", key=f"btn_{f_id}", icon=":material/lock:", disabled=True, use_container_width=True)
                    else:
                        st.markdown('<div style="height: 25px;"></div>', unsafe_allow_html=True) 
                        if st.button("進入分類", key=f"btn_{f_id}", icon=":material/login:", type="primary", use_container_width=True):
                            latest_doc = db.collection("workspaces").document(SHARED_WORKSPACE).get()
                            if latest_doc.exists:
                                latest_meta = latest_doc.to_dict().get("files_meta", [])
                                target_f = next((x for x in latest_meta if x["file_id"] == f_id), None)
                                
                                if target_f:
                                    cloud_locked_by = target_f.get("locked_by")
                                    cloud_locked_at = target_f.get("locked_at", 0)
                                    if cloud_locked_by and cloud_locked_by != current_uid and (time.time() - cloud_locked_at <= 180):
                                        st.toast("慢了一步！此檔案剛剛已被其他人進入。")
                                        load_from_cloud() 
                                        st.rerun()
                                        
                            f["locked_by"] = current_uid
                            f["locked_at"] = time.time()
                            st.session_state.active_file_id = f_id
                            sync_user_meta()
                            st.session_state.current_page = "題庫分類作業"
                            st.rerun()
                            
                with col4:
                    if is_admin:
                        st.markdown('<div style="height: 25px;"></div>', unsafe_allow_html=True) 
                        with st.popover("管理", icon=":material/settings:", use_container_width=True):
                            st.markdown("**修改考試資訊**")
                            new_year = st.number_input("考試年份", value=f.get("year", 114), key=f"y_{f_id}")
                            new_sess = st.number_input("考試次數", value=f.get("session", 2), key=f"s_{f_id}")
                            
                            if new_year != f.get("year") or new_sess != f.get("session"):
                                f["year"] = new_year
                                f["session"] = new_sess
                                
                                q_docs = list(db.collection("workspaces").document(SHARED_WORKSPACE).collection("questions").where("file_id", "==", f_id).stream())
                                batch = db.batch()
                                count = 0
                                for idx, doc in enumerate(q_docs):
                                    q_data = doc.to_dict()
                                    if q_data.get("category") != "未分類":
                                        new_year_info = f"{new_year}-{new_sess}-{q_data.get('order_index', idx) + 1}"
                                        batch.update(doc.reference, {"year_info": new_year_info}) 
                                        count += 1
                                        if count >= 400:
                                            batch.commit()
                                            batch = db.batch()
                                            count = 0
                                if count > 0: batch.commit()
                                sync_user_meta()
                                st.rerun()
                                
                            st.divider()
                            st.markdown("**重置分類**")
                            reset_check = st.text_input("請輸入 `Check` 以確認重置", key=f"check_reset_{f_id}")
                            if st.button("確認重置", key=f"reset_{f_id}", use_container_width=True):
                                if reset_check == "Check":
                                    q_docs = db.collection("workspaces").document(SHARED_WORKSPACE).collection("questions").where("file_id", "==", f_id).stream()
                                    batch = db.batch()
                                    for doc in q_docs:
                                        batch.update(doc.reference, {"category": "未分類", "categorizer_uid": None})
                                    batch.commit()
                                    
                                    f["categorized_qs"] = 0
                                    f["current_index"] = 0
                                    sync_user_meta()
                                    st.session_state.questions = []
                                    st.session_state.loaded_page = None
                                    st.success("已重置完成！")
                                    st.rerun()
                                else:
                                    st.error("輸入錯誤，請注意大小寫需為 Check", icon=":material/error:")
                                    
                            if st.button("刪除檔案", key=f"del_{f_id}", type="primary", use_container_width=True):
                                delete_file_data(f_id)
                                st.rerun()

        if auto_refresh:
            time.sleep(10)
            load_from_cloud()
            st.rerun()

    # ---------- 頁面 2：資料夾管理 ----------
    elif st.session_state.current_page == "資料夾管理":
        st.header("雲端資料夾管理")
        
        cat_counts = st.session_state.get("category_counts", {}) 

        folder_tree = {}
        for cat in st.session_state.categories:
            parts = cat.split("/")
            main_folder = parts[0]
            if main_folder not in folder_tree:
                folder_tree[main_folder] = []
            if len(parts) > 1:
                folder_tree[main_folder].append(cat)

        c_left, c_right = st.columns([1.5, 1])
        
        with c_left:
            st.markdown("### 瀏覽資料夾")
            with st.container(height=500, border=True):
                for main_folder, sub_folders in folder_tree.items():
                    main_qs_count = cat_counts.get(main_folder, 0)
                    total_qs_in_branch = main_qs_count + sum(cat_counts.get(sub, 0) for sub in sub_folders)
                    
                    is_expanded = main_folder in st.session_state.expanded_folders
                    icon_btn = ":material/folder_open:" if is_expanded else ":material/folder:"
                    
                    if st.button(f"{main_folder} (共 {total_qs_in_branch} 題)", key=f"toggle_{main_folder}", icon=icon_btn, use_container_width=True):
                        if is_expanded:
                            st.session_state.expanded_folders.remove(main_folder)
                        else:
                            st.session_state.expanded_folders.add(main_folder)
                        st.rerun()
                    
                    if is_expanded:
                        st.markdown(f"&nbsp;&nbsp;&nbsp;&nbsp;- **{main_folder} (根目錄)** : `{main_qs_count}` 題")
                        for sub in sub_folders:
                            sub_name = sub.split("/")[-1]
                            st.markdown(f"&nbsp;&nbsp;&nbsp;&nbsp;- {sub_name} : `{cat_counts.get(sub, 0)}` 題")
                            
        with c_right:
            st.markdown("### 新增與編輯")
            with st.container(border=True):
                c_main, c_sub = st.columns(2)
                main_folders = list(dict.fromkeys([c.split("/")[0] for c in st.session_state.categories]))
                
                parent_main = c_main.selectbox("新增位置 (母資料夾)", ["建立在主目錄 (最外層)"] + main_folders, key="add_new_main")
                if parent_main == "建立在主目錄 (最外層)":
                    parent_cat = "主目錄"
                else:
                    sub_options = [c for c in st.session_state.categories if c == parent_main or c.startswith(parent_main + "/")]
                    parent_cat = c_sub.selectbox("新增位置 (延伸子資料夾)", sub_options, key="add_new_sub")

                new_cat_name = st.text_input("新資料夾名稱 (多個用逗號分隔)")
                if st.button("新增", type="primary", use_container_width=True):
                    if new_cat_name:
                        for n in [x.strip() for x in new_cat_name.split(",") if x.strip()]:
                            full_path = n if "主目錄" in parent_cat else f"{parent_cat}/{n}"
                            if full_path not in st.session_state.categories: st.session_state.categories.append(full_path)
                        sync_user_meta()
                        st.rerun()

            with st.container(border=True):
                target_cat = hierarchical_select("選擇資料夾", st.session_state.categories, "edit_del")
                new_name = st.text_input("重新命名為", value=target_cat.split("/")[-1] if target_cat else "")
                c_e1, c_e2 = st.columns(2)
                if c_e1.button("儲存名稱", use_container_width=True):
                    old_name = target_cat.split("/")[-1]
                    if new_name != old_name:
                        base_path = "/".join(target_cat.split("/")[:-1])
                        new_full_path = f"{base_path}/{new_name}" if base_path else new_name
                        idx = st.session_state.categories.index(target_cat)
                        st.session_state.categories[idx] = new_full_path
                        
                        for i, c in enumerate(st.session_state.categories):
                            if c.startswith(target_cat + "/"):
                                st.session_state.categories[i] = c.replace(target_cat, new_full_path, 1)
                        for q in st.session_state.questions:
                            if q.get('category') == target_cat:
                                update_single_question_category(q["q_id"], new_full_path, q.get("year_info", ""))
                            elif q.get('category', '').startswith(target_cat + "/"):
                                update_single_question_category(q["q_id"], q['category'].replace(target_cat, new_full_path, 1), q.get("year_info", ""))
                        sync_user_meta()
                        st.rerun()
                        
                if c_e2.button("刪除", type="primary", use_container_width=True):
                    to_del = [c for c in st.session_state.categories if c == target_cat or c.startswith(target_cat + "/")]
                    for c in to_del: st.session_state.categories.remove(c)
                    
                    st.session_state.questions = []
                    st.session_state.loaded_page = None
                    sync_user_meta()
                    st.rerun()

            with st.container(border=True):
                st.markdown("**排序調整**")
                sort_level = st.radio("調整對象", ["母資料夾", "子資料夾"], horizontal=True, label_visibility="collapsed")
                
                if sort_level == "母資料夾":
                    main_folders = list(dict.fromkeys([c.split("/")[0] for c in st.session_state.categories]))
                    if main_folders:
                        sort_target = st.selectbox("選擇要移動的母資料夾", main_folders, key="sort_main")
                        
                        c_up, c_down = st.columns(2)
                        if c_up.button("上移", icon=":material/arrow_upward:", key="up_main", use_container_width=True):
                            idx = main_folders.index(sort_target)
                            if idx > 0:
                                main_folders[idx], main_folders[idx-1] = main_folders[idx-1], main_folders[idx]
                                new_categories = []
                                for mf in main_folders:
                                    new_categories.extend([c for c in st.session_state.categories if c.split("/")[0] == mf])
                                st.session_state.categories = new_categories
                                sync_user_meta()
                                st.rerun()
                                
                        if c_down.button("下移", icon=":material/arrow_downward:", key="down_main", use_container_width=True):
                            idx = main_folders.index(sort_target)
                            if idx < len(main_folders) - 1:
                                main_folders[idx], main_folders[idx+1] = main_folders[idx+1], main_folders[idx]
                                new_categories = []
                                for mf in main_folders:
                                    new_categories.extend([c for c in st.session_state.categories if c.split("/")[0] == mf])
                                st.session_state.categories = new_categories
                                sync_user_meta()
                                st.rerun()
                    else:
                        st.caption("尚無資料夾")
                        
                else:
                    main_folders = list(dict.fromkeys([c.split("/")[0] for c in st.session_state.categories]))
                    if main_folders:
                        parent_main = st.selectbox("選擇所屬的母資料夾", main_folders, key="sort_parent")
                        sub_folders = [c for c in st.session_state.categories if c.split("/")[0] == parent_main and c != parent_main]
                        
                        if not sub_folders:
                            st.caption("此母資料夾內無子資料夾")
                        else:
                            sort_target = st.selectbox("選擇要移動的子資料夾", sub_folders, key="sort_sub")
                            c_up, c_down = st.columns(2)
                            
                            if c_up.button("上移", icon=":material/arrow_upward:", key="up_sub", use_container_width=True):
                                idx = sub_folders.index(sort_target)
                                if idx > 0:
                                    sub_folders[idx], sub_folders[idx-1] = sub_folders[idx-1], sub_folders[idx]
                                    old_cat = st.session_state.categories.copy()
                                    positions = [i for i, c in enumerate(old_cat) if c in sub_folders]
                                    for pos, new_sub in zip(positions, sub_folders):
                                        old_cat[pos] = new_sub
                                    st.session_state.categories = old_cat
                                    sync_user_meta()
                                    st.rerun()
                                    
                            if c_down.button("下移", icon=":material/arrow_downward:", key="down_sub", use_container_width=True):
                                idx = sub_folders.index(sort_target)
                                if idx < len(sub_folders) - 1:
                                    sub_folders[idx], sub_folders[idx+1] = sub_folders[idx+1], sub_folders[idx]
                                    old_cat = st.session_state.categories.copy()
                                    positions = [i for i, c in enumerate(old_cat) if c in sub_folders]
                                    for pos, new_sub in zip(positions, sub_folders):
                                        old_cat[pos] = new_sub
                                    st.session_state.categories = old_cat
                                    sync_user_meta()
                                    st.rerun()
                    else:
                        st.caption("尚無資料夾")

    # ---------- 頁面 3：題庫分類作業 ----------
    elif st.session_state.current_page == "題庫分類作業":
        st.header("題庫分類作業")
        if not st.session_state.active_file_id:
            st.info("請先至「檔案管理」選擇要分類的檔案。", icon=":material/info:")
        else:
            active_f = next((f for f in st.session_state.files_meta if f["file_id"] == st.session_state.active_file_id), None)
            if not active_f: 
                st.error("找不到檔案", icon=":material/error:")
            else:
                # [優化] Lazy Loading: 只在進入時抓取該檔案的題目
                if st.session_state.loaded_page != f"classify_{active_f['file_id']}":
                    with st.spinner("載入試卷題目中..."):
                        raw_qs = [doc.to_dict() for doc in db.collection("workspaces").document(SHARED_WORKSPACE).collection("questions").where("file_id", "==", active_f["file_id"]).stream()]
                        st.session_state.questions = sorted(raw_qs, key=lambda x: x.get("order_index", 0))
                        st.session_state.loaded_page = f"classify_{active_f['file_id']}"
                        
                file_qs = st.session_state.questions
                if not file_qs:
                    st.warning("無有效題目", icon=":material/warning:")
                else:
                    curr_idx = active_f.get("current_index", 0)
                    total_q = len(file_qs)
                    
                    c_title, c_helper, c_exit = st.columns([3, 1, 1], vertical_alignment="center")
                    with c_title:
                        st.markdown(f"**檔案：{active_f['name']}**")
                        st.progress((curr_idx + 1) / total_q, text=f"進度：第 {curr_idx + 1} 題 / 共 {total_q} 題")
                    with c_helper:
                        # 懸浮中藥查詢按鈕
                        with st.popover("中藥查詢", icon=":material/lightbulb:", use_container_width=True):
                            st.markdown("**中藥分類查詢**")
                            search_kw = st.text_input("輸入關鍵字", key=f"herb_search_{active_f['file_id']}_{curr_idx}", placeholder="例如: 參")
                            if search_kw:
                                try:
                                    import cate
                                    import importlib
                                    importlib.reload(cate) # 強制重新載入，避免系統快取到舊的空白檔案
                                    
                                    results = cate.search_herb(search_kw.strip())
                                    if results:
                                        for r in results:
                                            # 將命中的關鍵字標記為紅色粗體
                                            herb_name = r['herb'].replace(search_kw.strip(), f":red[**{search_kw.strip()}**]")
                                            st.markdown(f"- {herb_name} : `{r['category']}`")
                                    else:
                                        st.caption("查無符合的中藥。")
                                except ImportError:
                                    st.error("找不到 cate.py 檔案", icon=":material/error:")
                                except AttributeError:
                                    st.error("請確認 cate.py 檔案已儲存，且包含 search_herb 函式。", icon=":material/error:")
                    with c_exit:
                        if st.button("退出分類", icon=":material/exit_to_app:", use_container_width=True):
                            active_f["locked_by"] = None
                            active_f["locked_at"] = None
                            st.session_state.active_file_id = None
                            sync_user_meta()
                            st.session_state.current_page = "檔案管理"
                            st.rerun()

                    current_q = file_qs[curr_idx]
                    
                    with st.container(border=True):
                        doubt_badge = " **(已標記為疑問)**" if current_q.get("is_doubt", False) else ""
                        st.caption(f"目前狀態： `{current_q.get('category', '未分類')}` {doubt_badge}")
                        for block in current_q.get('blocks', []):
                            if block['type'] == 'text': st.write(block['content'])
                            elif block['type'] == 'image': st.image(block['content'])
                        
                        render_categorizer_info(current_q.get("categorizer_uid"))

                    st.markdown("### 快速分類與標記")
                    if not st.session_state.categories:
                        st.info("尚無資料夾", icon=":material/info:")
                    else:
                        current_year_info = f"{active_f['year']}-{active_f['session']}-{curr_idx + 1}"
                        
                        is_doubt = st.checkbox("標記為疑問 (加入疑問區)", value=current_q.get("is_doubt", False), key=f"doubt_chk_{current_q['q_id']}")
                        if is_doubt != current_q.get("is_doubt", False):
                            update_single_question_category(current_q["q_id"], current_q.get("category", "未分類"), current_year_info, is_doubt)
                            active_f['locked_at'] = time.time()
                            sync_user_meta()

                        c_sel_main, c_sel_sub, c_btn = st.columns([1.5, 1.5, 1], vertical_alignment="bottom")
                        selected_cat = hierarchical_select(
                            "分類至", 
                            st.session_state.categories, 
                            "classify", 
                            default_cat=st.session_state.last_used_folder, 
                            col_layout=(c_sel_main, c_sel_sub)
                        )
                        
                        if c_btn.button("確定分類", type="primary", use_container_width=True, icon=":material/check_circle:"):
                            update_single_question_category(current_q["q_id"], selected_cat, current_year_info, is_doubt)
                            
                            st.session_state.last_used_folder = selected_cat
                            if selected_cat in st.session_state.recent_folders:
                                st.session_state.recent_folders.remove(selected_cat)
                            st.session_state.recent_folders.insert(0, selected_cat)
                            if len(st.session_state.recent_folders) > 5:
                                st.session_state.recent_folders = st.session_state.recent_folders[:5]
                                
                            if curr_idx < total_q - 1:
                                active_f['current_index'] = curr_idx + 1
                                
                            active_f['locked_at'] = time.time()
                            sync_user_meta()
                            st.rerun()

                        if st.session_state.recent_folders:
                            st.markdown("##### 最近使用的分類快捷鍵")
                            recent_cols = st.columns(5)
                            for i, r_cat in enumerate(st.session_state.recent_folders):
                                display_name = r_cat.split('/')[-1]
                                if recent_cols[i].button(display_name, key=f"recent_{r_cat}_{current_q['q_id']}", help=r_cat, use_container_width=True):
                                    update_single_question_category(current_q["q_id"], r_cat, current_year_info, is_doubt)
                                    
                                    st.session_state.last_used_folder = r_cat
                                    st.session_state.recent_folders.remove(r_cat)
                                    st.session_state.recent_folders.insert(0, r_cat)
                                    
                                    if curr_idx < total_q - 1:
                                        active_f['current_index'] = curr_idx + 1
                                        
                                    active_f['locked_at'] = time.time()
                                    sync_user_meta()
                                    st.rerun()

                    st.divider()
                    c1, c2, c3, c4 = st.columns(4)
                    if c1.button("上一題", icon=":material/arrow_back:") and curr_idx > 0:
                        active_f['current_index'] = curr_idx - 1; active_f['locked_at'] = time.time(); sync_user_meta(); st.rerun()
                    if c2.button("下一題", icon=":material/arrow_forward:") and curr_idx < total_q - 1:
                        active_f['current_index'] = curr_idx + 1; active_f['locked_at'] = time.time(); sync_user_meta(); st.rerun()
                    jump_to = c3.number_input("跳題", 1, total_q, curr_idx + 1, label_visibility="collapsed")
                    if c4.button("跳轉", icon=":material/keyboard_tab:"):
                        active_f['current_index'] = jump_to - 1; active_f['locked_at'] = time.time(); sync_user_meta(); st.rerun()

    # ---------- 頁面 4：全站題目搜索與瀏覽 ----------
    elif st.session_state.current_page == "全站題目搜索與瀏覽":
        st.header("全站題目搜索與瀏覽")
        
        with st.container(border=True):
            st.subheader("篩選條件")
            s_col1, s_col2, s_col3, s_col4 = st.columns([2, 1.5, 1.5, 1], vertical_alignment="bottom")
            search_text = s_col1.text_input("關鍵字搜尋 (針對題目內容)")
            
            main_folders = list(dict.fromkeys([c.split("/")[0] for c in st.session_state.categories]))
            main_options = ["所有分類", "未分類"] + main_folders
            selected_main = s_col2.selectbox("指定母資料夾", main_options, key="filter_main")

            if selected_main in ["所有分類", "未分類"]:
                filter_cat = selected_main
                s_col3.selectbox("指定子資料夾", ["-"], disabled=True, key="filter_sub")
            else:
                sub_options = ["(全部)"] + [c for c in st.session_state.categories if c == selected_main or c.startswith(selected_main + "/")]
                selected_sub = s_col3.selectbox("指定子資料夾", sub_options, key="filter_sub")
                filter_cat = selected_main if selected_sub == "(全部)" else selected_sub
            
            only_doubt = s_col4.checkbox("只看疑問區", value=False)
            
        # [優化] 依據搜尋條件 Lazy Loading 精準下載
        query_key = f"search_{filter_cat}_{only_doubt}_{search_text.strip()}"
        if st.session_state.loaded_page != query_key:
            with st.spinner("撈取雲端題目中..."):
                q_ref = db.collection("workspaces").document(SHARED_WORKSPACE).collection("questions")
                
                if filter_cat != "所有分類":
                    if filter_cat == "未分類":
                        docs = q_ref.where("category", "==", "未分類").stream()
                    else:
                        # Firestore 前綴搜尋法
                        docs = q_ref.where("category", ">=", filter_cat).where("category", "<", filter_cat + "\uf8ff").stream()
                elif only_doubt:
                    docs = q_ref.where("is_doubt", "==", True).stream()
                elif search_text.strip():
                    docs = q_ref.stream()
                else:
                    docs = q_ref.limit(100).stream()
                    st.info("提示：目前為預覽模式（顯示前 100 題）。請輸入關鍵字或選擇分類進行精確搜尋。", icon=":material/info:")
                    
                raw_qs = [doc.to_dict() for doc in docs]
                st.session_state.questions = sorted(raw_qs, key=lambda x: (x.get("file_id", ""), x.get("order_index", 0)))
                st.session_state.loaded_page = query_key
                
        filtered_qs = []
        for q in st.session_state.questions:
            if search_text and search_text.lower() not in q.get('_raw_text', '').lower():
                continue
            if only_doubt and not q.get('is_doubt', False):
                continue
            filtered_qs.append(q)
            
        st.write(f"符合條件共 **{len(filtered_qs)}** 題")
        
        PAGE_SIZE = 15
        total_pages = math.ceil(len(filtered_qs) / PAGE_SIZE) if filtered_qs else 1
        
        if len(filtered_qs) > 0:
            page_num = st.number_input("頁碼", min_value=1, max_value=total_pages, value=1)
            start_idx = (page_num - 1) * PAGE_SIZE
            end_idx = start_idx + PAGE_SIZE
            
            for q in filtered_qs[start_idx:end_idx]:
                raw_text = q.get('_raw_text', '').strip()
                raw_text_single_line = " ".join(raw_text.splitlines())
                
                snippet_len = 30
                snippet = raw_text_single_line[:snippet_len] + "..." if len(raw_text_single_line) > snippet_len else raw_text_single_line
                    
                expander_icon = ":material/help:" if q.get("is_doubt", False) else ":material/article:"
                expander_title = f"[{q.get('year_info', '未標記')}] {snippet} - 分類: {q.get('category', '未分類')}"
                
                with st.expander(expander_title, icon=expander_icon):
                    for block in q.get('blocks', []):
                        if block['type'] == 'text': 
                            content = block['content']
                            if search_text:
                                pattern = re.compile(re.escape(search_text), re.IGNORECASE)
                                content = pattern.sub(lambda m: f":red[**{m.group(0)}**]", content)
                            st.markdown(content)
                        elif block['type'] == 'image': st.image(block['content'])
                    
                    st.divider()
                    render_categorizer_info(q.get("categorizer_uid"))
                    
                    col_act1, col_act2, col_act3 = st.columns([1.2, 1.2, 2])
                    if col_act1.button("移回未分類", key=f"del_{q.get('q_id')}", icon=":material/delete:"):
                        update_single_question_category(q["q_id"], "未分類", q.get("year_info", ""))
                        st.rerun()
                        
                    doubt_btn_text = "取消疑問標記" if q.get("is_doubt", False) else "標記為疑問"
                    if col_act2.button(doubt_btn_text, key=f"toggle_doubt_{q.get('q_id')}", icon=":material/help_center:"):
                        new_doubt_status = not q.get("is_doubt", False)
                        update_single_question_category(q["q_id"], q.get("category", "未分類"), q.get("year_info", ""), new_doubt_status)
                        st.rerun()
                        
                    with col_act3.popover("重新分類", icon=":material/drive_file_move:"):
                        c_main_pop = st.container()
                        c_sub_pop = st.container()
                        new_cat = hierarchical_select("選擇新分類", st.session_state.categories, f"reclass_{q['q_id']}", col_layout=(c_main_pop, c_sub_pop))
                        if st.button("確定移動", key=f"btn_reclass_{q['q_id']}", icon=":material/check_circle:", type="primary", use_container_width=True):
                            update_single_question_category(q["q_id"], new_cat, q.get("year_info", ""), q.get("is_doubt", False))
                            st.rerun()

    # ---------- 頁面 5：題目匯出 ----------
    elif st.session_state.current_page == "題目匯出":
        st.header("客製化題庫匯出")
        
        with st.container(border=True):
            st.subheader("匯出範圍與過濾")
            st.write("請點選母資料夾展開後，勾選欲匯出的子分類：")
            
            export_cats = []
            main_folders = list(dict.fromkeys([c.split("/")[0] for c in st.session_state.categories]))
            
            cols = st.columns(3)
            for i, main in enumerate(main_folders):
                sub_folders = [c for c in st.session_state.categories if c == main or c.startswith(main + "/")]
                
                with cols[i % 3].expander(main, icon=":material/folder:"):
                    c_sa, c_da = st.columns(2)
                    if c_sa.button("全選", key=f"sa_{main}", use_container_width=True):
                        for sub in sub_folders:
                            st.session_state[f"export_chk_{sub}"] = True
                        st.rerun()
                    if c_da.button("清空", key=f"da_{main}", use_container_width=True):
                        for sub in sub_folders:
                            st.session_state[f"export_chk_{sub}"] = False
                        st.rerun()
                    
                    st.divider()
                    
                    for sub in sub_folders:
                        if f"export_chk_{sub}" not in st.session_state:
                            st.session_state[f"export_chk_{sub}"] = True
                            
                        is_checked = st.checkbox(sub, value=st.session_state[f"export_chk_{sub}"], key=f"dynamic_chk_{sub}")
                        st.session_state[f"export_chk_{sub}"] = is_checked
                        
                        if is_checked:
                            export_cats.append(sub)

            st.divider()
            c_f1, c_f2 = st.columns(2)
            min_year = c_f1.number_input("最小年份", value=90, min_value=1)
            max_year = c_f2.number_input("最大年份", value=150, max_value=300)
            
            st.divider()
            st.subheader("匯出格式設定")
            c_opts1, c_opts2, c_opts3 = st.columns(3)
            export_format = c_opts1.radio("匯出結構", ["合併為單一 Word 檔", "依資料夾分別打包 (ZIP)"])
            include_year = c_opts2.checkbox("匯出年份題號標記 (如: [114-2-1])", value=True)
            include_images = c_opts2.checkbox("匯出圖片", value=True)
            custom_title = c_opts3.text_input("單一合併檔主標題", "客製化分類題庫")
            
            if st.button("開始產生文件", type="primary", icon=":material/download:"):
                if not export_cats:
                    st.warning("請至少選擇一個資料夾", icon=":material/warning:")
                else:
                    with st.spinner("正在從雲端抓取指定資料，這可能需要一點時間..."):
                        
                        # [優化] 使用 in 批量查詢（每 30 個條件一批），大幅減少連線與請求次數
                        q_ref = db.collection("workspaces").document(SHARED_WORKSPACE).collection("questions")
                        export_qs = []
                        for i in range(0, len(export_cats), 30):
                            batch_cats = export_cats[i:i+30]
                            docs = q_ref.where("category", "in", batch_cats).stream()
                            export_qs.extend([doc.to_dict() for doc in docs])
                            
                        if export_format == "合併為單一 Word 檔":
                            doc = Document()
                            doc.add_heading(custom_title, 0)
                            
                            for cat in export_cats:
                                all_cat_qs = [q for q in export_qs if q.get('category') == cat]
                                cat_qs = [q for q in all_cat_qs if min_year <= get_year_from_info(q.get('year_info', '')) <= max_year]
                                
                                if cat_qs:
                                    doc.add_heading(f'{cat.replace("/", " - ")}', 1)
                                    for q in cat_qs:
                                        if include_year: doc.add_paragraph(f"[{q.get('year_info', '')}]")
                                        for block in q.get('blocks', []):
                                            if block['type'] == 'text': doc.add_paragraph(block['content'])
                                            elif block['type'] == 'image' and include_images:
                                                try:
                                                    img_res = requests.get(block['content'], timeout=10)
                                                    if img_res.status_code == 200: doc.add_picture(io.BytesIO(img_res.content), width=Inches(4))
                                                except:
                                                    doc.add_paragraph("[圖片載入失敗]")
                                        doc.add_paragraph("---")
                            doc_buffer = io.BytesIO()
                            doc.save(doc_buffer)
                            st.download_button("下載完整題庫 (Docx)", doc_buffer.getvalue(), f"{custom_title}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

                        else:
                            zip_buffer = io.BytesIO()
                            with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
                                for cat in export_cats:
                                    all_cat_qs = [q for q in export_qs if q.get('category') == cat]
                                    cat_qs = [q for q in all_cat_qs if min_year <= get_year_from_info(q.get('year_info', '')) <= max_year]
                                    
                                    if cat_qs:
                                        doc = Document()
                                        doc.add_heading(f'{cat.replace("/", "_")} 題庫', 0)
                                        for q in cat_qs:
                                            if include_year: doc.add_paragraph(f"[{q.get('year_info', '')}]")
                                            for block in q.get('blocks', []):
                                                if block['type'] == 'text': doc.add_paragraph(block['content'])
                                                elif block['type'] == 'image' and include_images:
                                                    try:
                                                        img_res = requests.get(block['content'], timeout=10)
                                                        if img_res.status_code == 200: doc.add_picture(io.BytesIO(img_res.content), width=Inches(4))
                                                    except: pass
                                            doc.add_paragraph("---")
                                        doc_buffer = io.BytesIO()
                                        doc.save(doc_buffer)
                                        zip_file.writestr(f"分類_{cat.replace('/', '_')}.docx", doc_buffer.getvalue())
                            st.download_button("下載壓縮檔 (ZIP)", zip_buffer.getvalue(), "分類題庫打包.zip", "application/zip")

    # ---------- 頁面 6：個人檔案 ----------
    elif st.session_state.current_page == "個人檔案":
        st.header("個人檔案")
        load_user_profile()
        load_all_users()
        
        user_prof = st.session_state.user_profile
        
        tabs = st.tabs(["我的檔案", "分類貢獻榜", "系統管理(管理員專用)"] if is_admin else ["我的檔案", "分類貢獻榜"])
        
        with tabs[0]:
            c_info1, c_info2 = st.columns([1, 2])
            with c_info1:
                if user_prof.get("avatar_url"):
                    st.markdown(f'<img src="{user_prof["avatar_url"]}" width="150" height="150" style="border-radius:50%; object-fit:cover;">', unsafe_allow_html=True)
                else:
                    st.info("尚未設定頭像", icon=":material/account_circle:")
                
                st.markdown(" ")
                selected_avatar_name = st.selectbox("選擇您的預設頭像", list(AVATAR_OPTIONS.keys()))
                if st.button("套用預設頭像", icon=":material/image:"):
                    new_url = AVATAR_OPTIONS[selected_avatar_name]
                    db.collection("users").document(st.session_state.user['localId']).update({"avatar_url": new_url})
                    st.toast("頭像已更新！")
                    st.rerun()

            with c_info2:
                new_name = st.text_input("個人暱稱", value=user_prof.get("nickname", ""))
                if st.button("儲存暱稱", icon=":material/save:"):
                    db.collection("users").document(st.session_state.user['localId']).update({"nickname": new_name})
                    st.toast("暱稱已更新")
                    st.rerun()

                st.divider()
                st.metric("已分類題數 (總計)", user_prof.get("categorized_count", 0))
                st.metric("目前擁有點數", user_prof.get("points", 0))

                if st.button("兌換星巴克飲料券 (需 500 點)", icon=":material/local_cafe:", type="primary", disabled=user_prof.get("points", 0) < 500):
                    new_pts = user_prof["points"] - 500
                    code = "".join(random.choices(string.ascii_uppercase + string.digits, k=14))
                    item = {"name": "星巴克飲料券", "code": code}
                    db.collection("users").document(st.session_state.user['localId']).update({
                        "points": new_pts,
                        "inventory": firestore.ArrayUnion([item])
                    })
                    st.success(f"兌換成功！已獲得隨機序號：{code}")
                    st.rerun()

            st.divider()
            st.subheader("我的物品欄")
            user_inv = user_prof.get("inventory", [])
            if not user_inv:
                st.caption("目前沒有任何物品")
            else:
                for item in user_inv:
                    st.markdown(f"- **{item['name']}** (兌換序號: `{item['code']}`)")

            st.divider()
            st.subheader("通知中心")
            user_noti = user_prof.get("notifications", [])
            if not user_noti:
                st.caption("目前沒有最新通知")
            else:
                for msg in reversed(user_noti):
                    st.markdown(f"- {msg}")

        with tabs[1]:
            st.subheader("分類貢獻榜")
            cols = st.columns(4)
            for idx, (uid, prof) in enumerate(st.session_state.all_users.items()):
                with cols[idx % 4].container(border=True):
                    if prof.get("avatar_url"):
                        st.markdown(f'<img src="{prof["avatar_url"]}" width="60" height="60" style="border-radius:50%; object-fit:cover;">', unsafe_allow_html=True)
                    else:
                        st.markdown(":material/account_circle:")
                    st.markdown(f"**{prof.get('nickname', '未知')}**")
                    st.caption(f"已分類: {prof.get('categorized_count', 0)} 題")

        if is_admin:
            with tabs[2]:
                st.subheader("管理員控制台")
                user_options = list(st.session_state.all_users.keys())
                target_uid = st.selectbox("選擇指定用戶", options=user_options, format_func=lambda x: f"{st.session_state.all_users[x].get('nickname', x)} ({st.session_state.all_users[x].get('email', '未知')})")

                if target_uid:
                    tgt_prof = st.session_state.all_users[target_uid]
                    st.write(f"目前點數: {tgt_prof.get('points', 0)}")
                    adj_pts = st.number_input("調整點數 (正數增加，負數減少)", value=0)
                    reason = st.text_input("調整原因 (必填)")
                    if st.button("確認調整點數", icon=":material/edit:"):
                        if reason:
                            new_pts = tgt_prof.get("points", 0) + adj_pts
                            msg = f"管理員已將您的點數調整 {adj_pts} 點。原因：{reason}"
                            db.collection("users").document(target_uid).update({
                                "points": new_pts,
                                "notifications": firestore.ArrayUnion([msg])
                            })
                            st.success("點數調整完成！已發送系統通知。")
                            st.rerun()
                        else:
                            st.warning("請填寫調整原因")

                    st.divider()
                    if st.button("強制重置該用戶所有分類", type="primary", icon=":material/warning:"):
                        batch = db.batch()
                        q_ref = db.collection("workspaces").document(SHARED_WORKSPACE).collection("questions")
                        docs = q_ref.where("categorizer_uid", "==", target_uid).stream()
                        
                        count = 0
                        for doc in docs:
                            batch.update(doc.reference, {"category": "未分類", "categorizer_uid": None})
                            count += 1
                            if count >= 400:
                                batch.commit()
                                batch = db.batch()
                                count = 0
                        if count > 0: 
                            batch.commit()

                        msg = f"系統管理員已強制重置您所分類的所有題目 (共 {count} 題)。"
                        db.collection("users").document(target_uid).update({
                            "notifications": firestore.ArrayUnion([msg])
                        })
                        st.success(f"已重置 {count} 題！已發送系統通知。")
                        st.rerun()

# ==========================================
# 自動登入與畫面路由
# ==========================================
if st.session_state.user is None and "remember_token" in st.query_params:
    try:
        token = st.query_params["remember_token"]
        user_info = auth.refresh(token)
        uid = user_info.get("user_id") or user_info.get("userId")
        
        if not uid:
            raise ValueError("無法從憑證取得有效的用戶 ID")
            
        user_info["localId"] = uid 
        
        full_user = admin_auth.get_user(uid)
        user_info["email"] = full_user.email
        
        st.session_state.user = user_info
        load_from_cloud()
        st.rerun() 
        
    except Exception as e:
        st.error(f"自動登入失效，請重新登入。系統訊息: {e}", icon=":material/warning:")
        del st.query_params["remember_token"]

if st.session_state.user is None:
    login_ui()
else:
    if "user_profile" not in st.session_state:
        load_user_profile()
        
    if not st.session_state.user_profile.get("is_setup_complete", True):
        setup_ui()
    else:
        main_app()
